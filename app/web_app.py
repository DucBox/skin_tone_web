import base64
import csv
import io
import os
import secrets
import mimetypes
import tempfile
from functools import lru_cache
from pathlib import Path

import cv2
from flask import Flask, render_template, request, redirect, url_for, send_file, session, abort
from werkzeug.exceptions import HTTPException

try:
    from app.history_store import HistoryStore
    from app.tool_calc_L import (
        SUPPORTED_IMAGE_EXTENSIONS,
        analyze_image_array,
        decode_image_bytes,
        normalize_brightness_scale,
    )
except ModuleNotFoundError:
    from history_store import HistoryStore
    from tool_calc_L import (
        SUPPORTED_IMAGE_EXTENSIONS,
        analyze_image_array,
        decode_image_bytes,
        normalize_brightness_scale,
    )


app = Flask(__name__, template_folder="templates")
app.secret_key = os.environ.get("SECRET_KEY") or secrets.token_hex(32)
app.config.update(SESSION_COOKIE_HTTPONLY=True, SESSION_COOKIE_SAMESITE="Lax")
DEFAULT_BRIGHTNESS_SCALE = normalize_brightness_scale(0.92)


@lru_cache(maxsize=1)
def history_store():
    return HistoryStore()


@app.context_processor
def csrf_context():
    if "csrf_token" not in session:
        session["csrf_token"] = secrets.token_urlsafe(32)
    return {"csrf_token": session["csrf_token"]}


@app.before_request
def check_csrf():
    if request.method == "POST":
        expected = session.get("csrf_token", "")
        supplied = request.form.get("csrf_token", "")
        if not expected or not secrets.compare_digest(expected, supplied):
            abort(400, description="Phiên làm việc đã hết hạn. Vui lòng tải lại trang rồi thử lại.")


def jpeg_bytes(image, max_size=None):
    if image is None:
        return None
    if max_size and max(image.shape[:2]) > max_size:
        ratio = max_size / max(image.shape[:2])
        image = cv2.resize(image, (max(1, round(image.shape[1] * ratio)),
                                   max(1, round(image.shape[0] * ratio))))
    ok, encoded = cv2.imencode(".jpg", image, [cv2.IMWRITE_JPEG_QUALITY, 95])
    if not ok:
        raise ValueError("Không mã hoá được ảnh.")
    return encoded.tobytes()


def decode_uploaded_file(file_storage):
    if file_storage is None or not file_storage.filename:
        return None

    file_bytes = file_storage.read()
    if not file_bytes:
        return None

    return decode_image_bytes(file_bytes)


def encode_image_data_url(image_bgr):
    ok, encoded = cv2.imencode(".jpg", image_bgr, [cv2.IMWRITE_JPEG_QUALITY, 95])
    if not ok:
        return None

    encoded_text = base64.b64encode(encoded.tobytes()).decode("ascii")
    return f"data:image/jpeg;base64,{encoded_text}"


def encode_csv_data_url(analyzed_items):
    buffer = io.StringIO()
    buffer.write("sep=;\n")
    writer = csv.DictWriter(
        buffer,
        fieldnames=[
            "duong_dan_anh",
            "trang_thai",
            "lstar_ma_trai",
            "lstar_ma_phai",
            "lstar_trung_binh",
            "nhom_mau_da",
            "mo_ta_nhom",
            "nguong_nhom",
            "thong_bao",
        ],
        delimiter=";",
        lineterminator="\n",
    )
    writer.writeheader()

    for item in analyzed_items:
        writer.writerow(
            {
                "duong_dan_anh": item["display_path"],
                "trang_thai": "Thanh cong" if item["ok"] else "Loi",
                "lstar_ma_trai": item["left_lstar"] if item["left_lstar"] is not None else "",
                "lstar_ma_phai": item["right_lstar"] if item["right_lstar"] is not None else "",
                "lstar_trung_binh": item["average_lstar"] if item["average_lstar"] is not None else "",
                "nhom_mau_da": item["skin_tone_group"] if item["skin_tone_group"] is not None else "",
                "mo_ta_nhom": item["skin_tone_label"] if item["skin_tone_label"] else "",
                "nguong_nhom": item["skin_tone_range"] if item["skin_tone_range"] else "",
                "thong_bao": item["message"],
            }
        )

    csv_bytes = ("\ufeff" + buffer.getvalue()).encode("utf-8")
    encoded_text = base64.b64encode(csv_bytes).decode("ascii")
    return f"data:text/csv;base64,{encoded_text}"


def allowed_image(filename):
    return Path(filename).suffix.lower() in SUPPORTED_IMAGE_EXTENSIONS


def build_item_result(display_path, analysis_result):
    item = {
        "display_path": display_path,
        "file_name": Path(display_path).name,
        "ok": analysis_result["ok"],
        "message": analysis_result["message"],
        "original_image_url": None,
        "visualization_image_url": None,
        "left_lstar": None,
        "right_lstar": None,
        "average_lstar": None,
        "skin_tone_group": None,
        "skin_tone_label": None,
        "skin_tone_range": None,
    }

    original_image = analysis_result.get("original_image")
    visualization_image = analysis_result.get("visualization_image")
    if original_image is not None:
        item["original_image_url"] = encode_image_data_url(original_image)
    if visualization_image is not None:
        item["visualization_image_url"] = encode_image_data_url(visualization_image)

    if analysis_result["ok"] and analysis_result["faces"]:
        first_face = analysis_result["faces"][0]
        skin_tone_group = first_face["skin_tone_group"]
        item.update(
            {
                "left_lstar": round(first_face["left_lstar"], 2),
                "right_lstar": round(first_face["right_lstar"], 2),
                "average_lstar": round(first_face["average_lstar"], 2),
                "skin_tone_group": skin_tone_group["group"] if skin_tone_group else None,
                "skin_tone_label": skin_tone_group["label"] if skin_tone_group else "Không có",
                "skin_tone_range": skin_tone_group["range_text"] if skin_tone_group else "Không có",
            }
        )

    return item


def analyze_files(file_storages, brightness_scale):
    results = []
    for file_storage in file_storages:
        if not file_storage or not file_storage.filename:
            continue
        if not allowed_image(file_storage.filename):
            continue

        original_bytes = file_storage.read()
        image = decode_image_bytes(original_bytes)
        if image is None:
            continue

        analysis_result = analyze_image_array(
            image,
            mode="largest",
            draw_labels=True,
            brightness_scale=brightness_scale,
        )
        item = build_item_result(file_storage.filename, analysis_result)
        original_type = mimetypes.guess_type(file_storage.filename)[0] or "application/octet-stream"
        if Path(file_storage.filename).suffix.lower() in (".heic", ".heif"):
            original_display = cv2.imencode(".png", image)[1].tobytes()
            display_type = "image/png"
        else:
            original_display = original_bytes
            display_type = original_type
        item["original_image_url"] = f"data:{display_type};base64," + base64.b64encode(original_display).decode("ascii")
        item["saved"] = False
        try:
            record = history_store().save(
                item, original_bytes,
                original_type,
                None,
                jpeg_bytes(analysis_result.get("visualization_image"), max_size=1600),
            )
            item["saved"] = True
            item["history_id"] = record["id"]
            item["original_image_url"] = url_for("history_image", record_id=record["id"], kind="preview")
            item["visualization_image_url"] = (
                url_for("history_image", record_id=record["id"], kind="visualization")
                if record["has_visualization"] else None)
        except Exception:
            app.logger.exception("Unable to save assessment history")
        results.append(item)

    return results


@app.route("/", methods=["GET", "POST"])
def index():
    analyzed_items = []
    error = None
    brightness_scale = DEFAULT_BRIGHTNESS_SCALE

    if request.method == "POST":
        single_image = request.files.get("single_image")
        folder_files = request.files.getlist("folder_files")

        if single_image and single_image.filename:
            analyzed_items = analyze_files([single_image], brightness_scale)
        elif folder_files:
            analyzed_items = analyze_files(folder_files, brightness_scale)

        if not analyzed_items:
            error = (
                "Vui lòng tải lên 1 ảnh hợp lệ hoặc chọn thư mục chứa ảnh "
                "(.png, .jpg, .jpeg, .webp, .heic, .heif)."
            )
        elif any(not item["saved"] for item in analyzed_items):
            error = "Có kết quả chưa lưu được vào lịch sử. Hãy kiểm tra kết nối lưu trữ và chấm lại các ảnh chưa lưu."

    selected_item = analyzed_items[0] if analyzed_items else None
    total_count = len(analyzed_items)
    success_count = sum(1 for item in analyzed_items if item["ok"])
    csv_export_url = encode_csv_data_url(analyzed_items) if analyzed_items else None

    return render_template(
        "index.html",
        analyzed_items=analyzed_items,
        selected_item=selected_item,
        error=error,
        total_count=total_count,
        success_count=success_count,
        csv_export_url=csv_export_url,
    )


@app.route("/history")
def history():
    error = None
    items, total, page = [], 0, max(1, request.args.get("page", 1, type=int))
    per_page = request.args.get("per_page", session.get("history_per_page", "20"))
    if per_page not in ("10", "20", "50", "100", "all"):
        per_page = "20"
    session["history_per_page"] = per_page
    page_count = 1
    try:
        store = history_store()
        ids = store.ids()
        total = len(ids)
        if per_page == "all":
            page = 1
            selected_ids = ids
        else:
            page_size = int(per_page)
            page_count = max(1, (total + page_size - 1) // page_size)
            page = min(page, page_count)
            selected_ids = ids[(page - 1) * page_size:page * page_size]
        for record_id in selected_ids:
            try:
                items.append(store.get(record_id))
            except FileNotFoundError:
                continue  # Another user may have just deleted this record.
    except Exception:
        app.logger.exception("Unable to read history")
        error = "Không tải được lịch sử. Vui lòng thử lại hoặc kiểm tra cấu hình lưu trữ."
    return render_template("history.html", items=items, total=total, page=page,
                           page_count=page_count, per_page=per_page, error=error)


@app.route("/history/<record_id>/image/<kind>")
def history_image(record_id, kind):
    names = {"original": "original", "preview": "original", "visualization": "visualization.jpg"}
    if kind not in names:
        abort(404)
    try:
        store = history_store()
        record = store.get(record_id)
        if kind == "visualization" and not record["has_visualization"]:
            abort(404)
        if kind == "preview" and Path(record["file_name"]).suffix.lower() in (".heic", ".heif"):
            # Browsers cannot generally display HEIC. Decode at full resolution,
            # encode losslessly to PNG for viewing; stored original is untouched.
            image = decode_image_bytes(store.read(store.key(record_id, "original")))
            if image is None:
                abort(404)
            ok, encoded = cv2.imencode(".png", image)
            if not ok:
                abort(503)
            return send_file(io.BytesIO(encoded.tobytes()), mimetype="image/png", as_attachment=False)
        if store.client:
            response = redirect(store.signed_url(record_id, names[kind]))
            response.headers["Cache-Control"] = "no-store"
            return response
        return send_file(io.BytesIO(store.read(store.key(record_id, names[kind]))),
                         mimetype=record["original_type"] if kind in ("original", "preview") else "image/jpeg",
                         download_name=record["file_name"], as_attachment=False)
    except (ValueError, FileNotFoundError):
        abort(404)
    except HTTPException:
        raise
    except Exception:
        app.logger.exception("Unable to read history image")
        abort(503)


@app.route("/history/<record_id>/delete", methods=["POST"])
def delete_history(record_id):
    try:
        history_store().delete(record_id)
    except FileNotFoundError:
        pass  # Deleting an already deleted row is harmless.
    except ValueError:
        abort(404)
    except Exception:
        app.logger.exception("Unable to delete history")
        return render_template("history.html", items=[], total=0, page=1, page_count=1,
                               error="Chưa xoá được đầy đủ kết quả. Vui lòng tải lại lịch sử và thử lại."), 503
    per_page = request.form.get("per_page", session.get("history_per_page", "20"))
    if per_page not in ("10", "20", "50", "100", "all"):
        per_page = "20"
    return redirect(url_for("history", page=max(1, request.form.get("page", 1, type=int)), per_page=per_page))


@app.route("/history/export.xlsx")
def export_history():
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as ExcelImage
    from openpyxl.styles import Font, PatternFill, Alignment
    from PIL import Image, ImageOps

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Lịch sử chấm"
    sheet.append(["Ảnh gốc", "Điểm chấm (L*)", "Khoảng L*", "Kết quả màu da"])
    sheet.freeze_panes = "B2"
    sheet.column_dimensions["A"].width = 30
    sheet.column_dimensions["B"].width = 23
    sheet.column_dimensions["C"].width = 22
    sheet.column_dimensions["D"].width = 40
    sheet.row_dimensions[1].height = 32
    for cell in sheet[1]:
        cell.font = Font(name="Calibri", bold=True, color="FFFFFF", size=12)
        cell.fill = PatternFill("solid", fgColor="325B84")
        cell.alignment = Alignment(vertical="center")
    try:
        store = history_store()
        scratch = tempfile.TemporaryDirectory(prefix="skin-tone-export-")
        for record_id in store.ids():
            try:
                item = store.get(record_id)
            except FileNotFoundError:
                continue
            # Convert HEIC/WebP/etc. to PNG at original resolution for Excel compatibility.
            original_bytes = store.read(store.key(record_id, "original"))
            source = io.BytesIO(original_bytes)
            with Image.open(source) as photo:
                if photo.format in ("JPEG", "PNG") and photo.getexif().get(274, 1) == 1:
                    embedded = Path(scratch.name) / f"{record_id}.{photo.format.lower()}"
                    embedded.write_bytes(original_bytes)
                else:
                    embedded = Path(scratch.name) / f"{record_id}.png"
                    ImageOps.exif_transpose(photo).convert("RGB").save(embedded, format="PNG")
            image = ExcelImage(str(embedded))
            ratio = min(200 / image.width, 150 / image.height, 1)
            image.width, image.height = image.width * ratio, image.height * ratio
            label = (f"Nhóm {item['skin_tone_group']} — {item['skin_tone_label']}"
                     if item["ok"] else f"Không chấm được: {item['message']}")
            sheet.append([item["file_name"], item["average_lstar"], item["skin_tone_range"] or "—", label])
            row = sheet.max_row
            sheet.add_image(image, f"A{row}")
            sheet.row_dimensions[row].height = 140
            for cell in sheet[row]:
                cell.font = Font(name="Calibri", size=11)
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                if row % 2 == 0:
                    cell.fill = PatternFill("solid", fgColor="F2F6FA")
            sheet.cell(row, 1).alignment = Alignment(vertical="bottom", wrap_text=True)
            sheet.cell(row, 1).data_type = "s"  # Uploaded filenames must never become formulas.
            sheet.cell(row, 2).number_format = "0.00"
        sheet.auto_filter.ref = f"A1:D{sheet.max_row}"
        output = tempfile.SpooledTemporaryFile(max_size=8 * 1024 * 1024)
        workbook.save(output)
        output.seek(0)
        return send_file(output, as_attachment=True, download_name="lich_su_cham_mau_da.xlsx",
                         mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    except Exception:
        app.logger.exception("Unable to export history")
        return render_template("history.html", items=[], total=0, page=1, page_count=1,
                               error="Không xuất được Excel. Vui lòng thử lại; lịch sử vẫn được giữ nguyên."), 503
    finally:
        if "scratch" in locals():
            scratch.cleanup()


@app.route("/history/delete-all", methods=["POST"])
def delete_all_history():
    try:
        history_store().delete_all()
    except Exception:
        app.logger.exception("Unable to delete all history")
        return render_template("history.html", items=[], total=0, page=1, page_count=1,
                               error="Chưa xoá được toàn bộ lịch sử. Vui lòng tải lại trang và thử lại."), 503
    return redirect(url_for("history"))


if __name__ == "__main__":
    host = os.environ.get("HOST", "127.0.0.1")
    port = int(os.environ.get("PORT", "5001"))
    debug = os.environ.get("FLASK_DEBUG", "1") == "1"
    app.run(host=host, port=port, debug=debug)
